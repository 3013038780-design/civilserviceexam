"""使用选定工作区Python运行，将完成的八课Markdown生成Word和网站下载包。"""
from pathlib import Path
import re,json,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parents[1]
M=ROOT/'manuscripts'
OUT=ROOT/'dist/downloads'
OUT.mkdir(parents=True,exist_ok=True)
def hyperlink(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'334B3E');pr.append(co);r.append(pr);t=OxmlElement('w:t');t.text=label+' ↗';r.append(t);el.append(r);p._p.append(el)
def inline(p,text):
    pattern=re.compile(r'\[([^\]]+)\]\((https?://[^)]+)\)')
    pos=0
    for m in pattern.finditer(text):
        if m.start()>pos:p.add_run(text[pos:m.start()])
        hyperlink(p,m.group(1),m.group(2));pos=m.end()
    if pos<len(text):p.add_run(text[pos:])
def setup():
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.82);sec.right_margin=Inches(.82)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name];st.font.name='Source Han Serif CN';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Source Han Serif CN');st.font.color.rgb=RGBColor(0,0,0)
    normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(4)
    for name,size in [('Title',23),('Heading 1',18),('Heading 2',13),('Heading 3',11.5)]:
        st=d.styles[name];st.font.size=Pt(size);st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(9);st.paragraph_format.space_after=Pt(4)
    h=sec.header.paragraphs[0];h.text='了解中国与申论八课讲义';h.runs[0].font.size=Pt(9);h.runs[0].font.color.rgb=RGBColor(0,0,0)
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('2026年10月17日  ·  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    for run in f.runs:run.font.size=Pt(9)
    
    for st in d.styles:
        for border in list(st._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for para in list(d.paragraphs)+list(sec.header.paragraphs)+list(sec.footer.paragraphs):
        for border in list(para._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.core_properties.author='公考备考协作项目';d.core_properties.subject='官方依据与原创申论教学';return d
def append(d,text,first=False,start_new_page=False):
    for line in text.splitlines():
        line=line.strip()
        if not line:continue
        if line.startswith('# '):
            p=d.add_paragraph(line[2:],style='Title' if first else 'Heading 1');first=False
            if start_new_page:p.paragraph_format.page_break_before=True;start_new_page=False
        elif line.startswith('## '):d.add_paragraph(line[3:],style='Heading 2')
        elif line.startswith('### '):d.add_paragraph(line[4:],style='Heading 3')
        else:
            p=d.add_paragraph();inline(p,line[2:] if line.startswith('- ') else line)
            if line.startswith('- '):p.paragraph_format.left_indent=Inches(.12)
            p.paragraph_format.widow_control=True
manifest=json.loads((M/'manifest.json').read_text())
for x in manifest:
    f=M/x['file'];d=setup();append(d,f.read_text(),first=True);d.save(OUT/(f.stem+'.docx'))
# 合订版显式分课，避免单课标题出现在前一课段尾。
d=setup();append(d,'# 了解中国与申论八课讲义\n\n从官方依据到实际工作再到独立表达\n\n本版为八节完整讲义初稿。十九章主题均有对应，实际选读范围见章节对照。先理解，再记忆必要概念；讲义不替代题干材料。\n\n版本日期 2026年10月17日\n\n## 八课目录\n\n'+'\n'.join('- 第'+str(x['unit'])+'课 '+x['title'] for x in manifest),first=True)
append(d,(M/'00-章节与来源对照.md').read_text(),start_new_page=True)
for x in manifest:append(d,(M/x['file']).read_text(),start_new_page=True)
d.save(OUT/'了解中国与申论八课讲义.docx')
# 供网页和直接下载使用的Markdown保持与源稿一致。
for f in M.glob('*.md'):(OUT/f.name).write_bytes(f.read_bytes())
with zipfile.ZipFile(OUT/'八课Markdown与Word.zip','w',zipfile.ZIP_DEFLATED) as z:
    for f in sorted(OUT.iterdir()):
        if f.suffix in ['.md','.docx']:z.write(f,f.name)
print(json.dumps({'docx':len(list(OUT.glob('*.docx'))),'markdown':len(list(OUT.glob('*.md'))),'output':str(OUT)},ensure_ascii=False))
